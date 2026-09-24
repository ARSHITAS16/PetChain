import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Colors
    NAVY_DARK = RGBColor(11, 25, 44)       # #0B192C
    NAVY_MAIN = RGBColor(15, 44, 89)       # #0F2C59
    BLUE_ACCENT = RGBColor(30, 64, 175)     # #1E40AF
    CYAN_PRIMARY = RGBColor(2, 132, 199)    # #0284C7
    TEXT_DARK = RGBColor(31, 41, 55)       # #1F2937
    TEXT_MUTED = RGBColor(107, 114, 128)   # #6B7280
    BG_LIGHT = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(243, 244, 246)      # #F3F4F6
    CARD_BORDER = RGBColor(229, 231, 235)  # #E5E7EB
    GREEN_SUCCESS = RGBColor(16, 185, 129)  # #10B981
    YELLOW_NOTE = RGBColor(254, 240, 138)   # #FEF08A

    blank_layout = prs.slide_layouts[6]

    # Helper: Add Header Bar to Content Slides (Slides 3 to 7)
    def add_slide_header(slide, title_text):
        # Title text
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = 'Arial'
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = NAVY_MAIN

        # Solid Underline Rule
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = NAVY_MAIN
        line.line.fill.background()

    # Helper: Add Bottom Presidency Banner (Slides 2 to 8)
    def add_bottom_banner(slide):
        banner = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(-0.5), Inches(6.5), Inches(14.333), Inches(1.2))
        banner.rotation = 180
        banner.fill.solid()
        banner.fill.fore_color.rgb = NAVY_MAIN
        banner.line.fill.background()

        # Banner Text / Logo Simulation
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.75), Inches(10), Inches(0.6))
        tf = tx_box.text_frame
        p = tf.paragraphs[0]
        p.text = "PRESIDENCY UNIVERSITY  |  School of Computer Science and Engineering"
        p.font.name = 'Arial'
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

    # -------------------------------------------------------------------------
    # SLIDE 1: Cover Title
    # -------------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_DARK
    bg1.line.fill.background()

    # Header branding
    tx = slide1.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(6), Inches(0.5))
    p = tx.text_frame.paragraphs[0]
    p.text = "PRESIDENCY UNIVERSITY"
    p.font.name = 'Arial'
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Slide number
    tx_num = slide1.shapes.add_textbox(Inches(12.0), Inches(0.5), Inches(0.8), Inches(0.5))
    p_num = tx_num.text_frame.paragraphs[0]
    p_num.text = "1"
    p_num.alignment = PP_ALIGN.RIGHT
    p_num.font.name = 'Arial'
    p_num.font.size = Pt(14)
    p_num.font.color.rgb = RGBColor(200, 200, 200)

    # Big Center Text
    tx_center = slide1.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.333), Inches(2.5))
    tf_c = tx_center.text_frame
    tf_c.word_wrap = True
    
    p1 = tf_c.paragraphs[0]
    p1.text = "Presidency"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = 'Arial'
    p1.font.size = Pt(72)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)

    p2 = tf_c.add_paragraph()
    p2.text = "School of Computer Science and Engineering"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = 'Arial'
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(186, 230, 253)

    # -------------------------------------------------------------------------
    # SLIDE 2: Title & Project Info Review
    # -------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    
    # Sub-header
    tx2_sub = slide2.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.6))
    p = tx2_sub.text_frame.paragraphs[0]
    p.text = "CSSE2289-Foundation Of Blockchain Technology Mini Project Synopsis Review"
    p.font.name = 'Arial'
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_MAIN
    p.alignment = PP_ALIGN.CENTER

    line2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.0), Inches(11.733), Inches(0.04))
    line2.fill.solid()
    line2.fill.fore_color.rgb = NAVY_MAIN
    line2.line.fill.background()

    # Main Title
    tx2_title = slide2.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(1.2))
    p = tx2_title.text_frame.paragraphs[0]
    p.text = "“PetChain: A Blockchain-Based Pet Adoption and Ownership Management System”"
    p.font.name = 'Arial'
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK
    p.alignment = PP_ALIGN.CENTER

    # Student Table / Info
    info_box = slide2.shapes.add_textbox(Inches(1.5), Inches(3.0), Inches(10.333), Inches(1.2))
    tf = info_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Roll Number\t\t\tStudent Name"
    p.font.name = 'Arial'
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY_MAIN

    p = tf.add_paragraph()
    p.text = "20231CSE0001\t\t\tTeam Member 1 (Lead)\n20231CSE0002\t\t\tTeam Member 2\n20231CSE0003\t\t\tTeam Member 3"
    p.font.name = 'Arial'
    p.font.size = Pt(16)
    p.font.color.rgb = TEXT_DARK

    # Guide Info
    guide_box = slide2.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(10), Inches(1.5))
    tf_g = guide_box.text_frame
    
    g1 = tf_g.paragraphs[0]
    g1.text = "Project Guide :- Mr. Akash Raj (Assistant Professor)"
    g1.font.name = 'Arial'
    g1.font.size = Pt(18)
    g1.font.bold = True
    g1.font.color.rgb = CYAN_PRIMARY

    g2 = tf_g.add_paragraph()
    g2.text = "Name of the Program: B.tech\nName of Branch:- Computer Science and Engineering"
    g2.font.name = 'Arial'
    g2.font.size = Pt(16)
    g2.font.bold = True
    g2.font.color.rgb = CYAN_PRIMARY

    add_bottom_banner(slide2)

    # -------------------------------------------------------------------------
    # SLIDE 3: Problem Statement
    # -------------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide3, "Problem Statement")

    problems = [
        ("❌ Lack of Ownership Transparency", "Traditional paper-based adoption certificates can be lost, altered, or forged, leading to ownership disputes and illegal re-homing."),
        ("❌ Unverified Medical & Vaccination Logs", "Vaccination records are easily falsified or misplaced, risking animal health and giving adopters zero verifiable health proof."),
        ("❌ Centralized & Isolated Data Silos", "Shelters operate fragmented, offline databases prone to data loss, corruption, and complete lack of public auditability."),
        ("❌ Absence of Traceable Pet History", "Adopters have no immutable mechanism to verify a pet's complete lifecycle timeline, prior shelter transfers, or treatment history.")
    ]

    top_pos = 1.4
    for title, desc in problems:
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(11.733), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = 'Arial'
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = RGBColor(220, 38, 38) # Red accent

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = 'Arial'
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_DARK

        top_pos += 1.25

    add_bottom_banner(slide3)

    # -------------------------------------------------------------------------
    # SLIDE 4: Proposed Solution
    # -------------------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide4, "Proposed Solution")

    solutions = [
        ("🐾 Decentralized Pet Registry", "Ethereum smart contract (`PetChain.sol`) serving as an immutable, tamper-proof single source of truth for pet identity.", CYAN_PRIMARY),
        ("📜 Immutable Ownership History", "Cryptographically signed timeline logging every ownership transition (Shelter ➔ Adopter A ➔ Adopter B) permanently on-chain.", NAVY_MAIN),
        ("💉 Verified Medical & Vaccination Logs", "Authorized shelter administrators log vaccination dates, vaccine types, and clinic details directly on-chain.", GREEN_SUCCESS),
        ("🤖 Smart Contract Automation", "Decentralized adoption lifecycle (Request ➔ Review ➔ Approve/Reject) with automated competing request cancellation.", BLUE_ACCENT),
        ("🦊 Web3 & MetaMask Integration", "Seamless browser interaction via `ethers.js` v6, Web3 wallet signatures, and real-time transaction hash tracking.", CYAN_PRIMARY)
    ]

    top_pos = 1.35
    for title, desc, color in solutions:
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(11.733), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.12)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = 'Arial'
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = 'Arial'
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

        top_pos += 1.05

    add_bottom_banner(slide4)

    # -------------------------------------------------------------------------
    # SLIDE 5: How the System Works (Flowchart + Step-by-Step)
    # -------------------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide5, "How the System Works")

    # Flowchart Visual Boxes
    flow_steps = [
        ("1. Registration 📝", "Shelter registers pet\non-chain (PetChain.sol)"),
        ("2. Marketplace 🐾", "Adopter views pets\nsynced via ethers.js"),
        ("3. Application 📩", "User submits adoption\nrequest & notes"),
        ("4. Admin Review 🛡️", "Shelter approves or\nrejects application"),
        ("5. Transfer & Logs 📜", "Ownership transferred,\nvaccination logged")
    ]

    box_width = 2.1
    gap = 0.25
    left_start = 0.8

    for idx, (title, desc) in enumerate(flow_steps):
        cur_left = left_start + idx * (box_width + gap)

        # Flow Box
        box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cur_left), Inches(1.4), Inches(box_width), Inches(1.5))
        box.fill.solid()
        box.fill.fore_color.rgb = NAVY_MAIN
        box.line.color.rgb = CYAN_PRIMARY

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.1)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = 'Arial'
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = 'Arial'
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(224, 242, 254)
        p2.alignment = PP_ALIGN.CENTER

        # Connector Arrow (except for last item)
        if idx < 4:
            arrow = slide5.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cur_left + box_width + 0.05), Inches(2.0), Inches(0.15), Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = CYAN_PRIMARY
            arrow.line.fill.background()

    # Workflow Detailed Explanation Box below Flowchart
    detail_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.1), Inches(11.733), Inches(3.2))
    detail_box.fill.solid()
    detail_box.fill.fore_color.rgb = CARD_BG
    detail_box.line.color.rgb = CARD_BORDER

    tf_d = detail_box.text_frame
    tf_d.word_wrap = True
    tf_d.margin_left = Inches(0.25)
    tf_d.margin_top = Inches(0.2)

    p = tf_d.paragraphs[0]
    p.text = "Workflow Architecture & State Lifecycle"
    p.font.name = 'Arial'
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY_MAIN

    bullets = [
        "🔹 User / Adopter Workflow: Connects Web3 Wallet ➔ Browses live pet gallery ➔ Submits adoption request with housing notes ➔ Receives instant transaction confirmation.",
        "🔹 Shelter Admin Workflow: Accesses Admin Dashboard ➔ Registers new pet profile ➔ Reviews pending applications ➔ Approves adoption (triggering atomic ownership transfer) ➔ Logs medical vaccination records.",
        "🔹 Smart Contract Automation: On approval, `PetChain.sol` updates owner address, sets `isAdopted = true`, appends `OwnershipRecord`, and automatically marks competing requests as `Rejected`."
    ]

    for b in bullets:
        p_b = tf_d.add_paragraph()
        p_b.text = b
        p_b.font.name = 'Arial'
        p_b.font.size = Pt(13)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_before = Pt(6)

    add_bottom_banner(slide5)

    # -------------------------------------------------------------------------
    # SLIDE 6: Requirements
    # -------------------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide6, "Requirements")

    # Column 1: Software Requirements
    col1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(3.6), Inches(4.8))
    col1.fill.solid()
    col1.fill.fore_color.rgb = CARD_BG
    col1.line.color.rgb = CARD_BORDER

    tf1 = col1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.2)
    tf1.margin_top = Inches(0.2)

    p = tf1.paragraphs[0]
    p.text = "1) Software Requirements"
    p.font.name = 'Arial'
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY_MAIN

    sw_items = [
        "• Smart Contract: Solidity 0.8.24",
        "• Framework: Hardhat 2.22",
        "• Frontend: React 18, Vite 5",
        "• Web3 Middleware: ethers.js v6",
        "• Wallet: MetaMask Extension",
        "• Runtime: Node.js v18+ / v24+",
        "• Environment: Windows / Linux"
    ]
    for item in sw_items:
        p_item = tf1.add_paragraph()
        p_item.text = item
        p_item.font.name = 'Arial'
        p_item.font.size = Pt(13)
        p_item.font.color.rgb = TEXT_DARK
        p_item.space_before = Pt(8)

    # Column 2: Hardware Requirements
    col2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.85), Inches(1.4), Inches(3.6), Inches(4.8))
    col2.fill.solid()
    col2.fill.fore_color.rgb = CARD_BG
    col2.line.color.rgb = CARD_BORDER

    tf2 = col2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.2)
    tf2.margin_top = Inches(0.2)

    p = tf2.paragraphs[0]
    p.text = "2) Hardware Requirements"
    p.font.name = 'Arial'
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY_MAIN

    hw_items = [
        "• Processor: Dual-Core CPU 2.0 GHz+",
        "• System RAM: 8 GB minimum",
        "• Storage: 10 GB SSD space",
        "• Network: Localhost / Internet",
        "• Display: 1280 x 720 minimum",
        "• Input: Keyboard & Mouse"
    ]
    for item in hw_items:
        p_item = tf2.add_paragraph()
        p_item.text = item
        p_item.font.name = 'Arial'
        p_item.font.size = Pt(13)
        p_item.font.color.rgb = TEXT_DARK
        p_item.space_before = Pt(8)

    # Column 3: Blockchain Network Specs
    col3 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(1.4), Inches(3.6), Inches(4.8))
    col3.fill.solid()
    col3.fill.fore_color.rgb = CARD_BG
    col3.line.color.rgb = CARD_BORDER

    tf3 = col3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = Inches(0.2)
    tf3.margin_top = Inches(0.2)

    p = tf3.paragraphs[0]
    p.text = "3) Blockchain Network Specs"
    p.font.name = 'Arial'
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY_MAIN

    bc_items = [
        "• Network: Hardhat Localhost",
        "• Chain ID: 31337 (0x7a69)",
        "• RPC URL: http://127.0.0.1:8545",
        "• Currency: Test ETH",
        "• Contract Address: 0x5FbD...80aa3",
        "• Unit Tests: 17/17 Passed (100%)"
    ]
    for item in bc_items:
        p_item = tf3.add_paragraph()
        p_item.text = item
        p_item.font.name = 'Arial'
        p_item.font.size = Pt(13)
        p_item.font.color.rgb = TEXT_DARK
        p_item.space_before = Pt(8)

    add_bottom_banner(slide6)

    # -------------------------------------------------------------------------
    # SLIDE 7: Expected Outcome
    # -------------------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide7, "Expected Outcome")

    outcomes = [
        ("✅ 100% Transparent & Verifiable Pet Records", "Publicly accessible Web3 dApp enabling adopters and shelters to inspect pet profiles without central authority trust."),
        ("✅ Tamper-Proof Ownership Transfer", "Guaranteed proof of ownership with an immutable, timestamped timeline recording every transition from shelter to adopters."),
        ("✅ Immutable Medical & Vaccination History", "Eliminates fraudulent health certificates by storing verified vaccination logs directly on-chain."),
        ("✅ Automated & Secure Adoption Management", "Smart contract driven approval/rejection lifecycle that auto-rejects competing applications upon adoption finalization."),
        ("✅ Fully Functioning Web3 Prototype", "Successfully deployed and tested on local EVM with 17 passing Hardhat unit tests and live Web3 frontend.")
    ]

    top_pos = 1.35
    for title, desc in outcomes:
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(11.733), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.12)

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = 'Arial'
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = GREEN_SUCCESS

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = 'Arial'
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

        top_pos += 1.05

    add_bottom_banner(slide7)

    # -------------------------------------------------------------------------
    # SLIDE 8: Thank You!
    # -------------------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)

    # Underline Header Bar
    line8 = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.04))
    line8.fill.solid()
    line8.fill.fore_color.rgb = NAVY_MAIN
    line8.line.fill.background()

    # Sticky Note Box
    note = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.6), Inches(2.2), Inches(4.2), Inches(3.2))
    note.fill.solid()
    note.fill.fore_color.rgb = YELLOW_NOTE
    note.line.color.rgb = RGBColor(250, 204, 21) # Darker yellow border
    note.rotation = -2 # Subtle realistic tilt

    # Blue Pin on top of sticky note
    pin = slide8.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.5), Inches(2.05), Inches(0.35), Inches(0.35))
    pin.fill.solid()
    pin.fill.fore_color.rgb = RGBColor(30, 64, 175) # Blue pin
    pin.line.fill.background()

    tf_n = note.text_frame
    tf_n.word_wrap = True
    
    p = tf_n.paragraphs[0]
    p.text = "\nThank\nYou!"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = 'Comic Sans MS' # Matches handwritten style on sticky note
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(17, 24, 39)

    add_bottom_banner(slide8)

    # Save presentation
    output_path = os.path.join(os.path.dirname(__file__), "PetChain_Presentation.pptx")
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {output_path}")

if __name__ == '__main__':
    build_presentation()
